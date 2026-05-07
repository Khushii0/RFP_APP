import os
import json
import PyPDF2
import docx
from pptx import Presentation
import openpyxl
from openai import OpenAI
from dotenv import load_dotenv

# Load env variables from the agent folder
load_dotenv(r"c:\Users\adarsh.s\Music\rfp_agent3 - new\.env")
client = OpenAI(api_key=os.environ.get("OPENAI_API_KEY"))

rfp_dir = "RFP collection"
target_json = "rfp_roadmap_strategy_full.json"

# Load existing JSON
try:
    with open(target_json, "r", encoding="utf-8") as f:
        data = json.load(f)
        if isinstance(data, dict):
            all_rfps = [data]
        elif isinstance(data, list):
            all_rfps = data
        else:
            all_rfps = []
except Exception as e:
    all_rfps = []

processed_files = set(rfp.get("_source_file") for rfp in all_rfps if "_source_file" in rfp)

# Adani example
example_json = """
{
    "industry_vertical_context": "The project is for Adani Group, a large conglomerate...",
    "high_level_strategic_objective": "To design, develop, and implement...",
    "descriptive_pain_points": [
        "Lack of Price Transparency...",
        "Fragmented Procurement Process..."
    ],
    "horizontal_capabilities": [
        {
            "category": "Horizontal",
            "feature_name": "Unified Search Interface",
            "functional_specification": "...",
            "non_functional_constraints": "...",
            "source_filename": "example.pdf",
            "section_header": "4.1"
        }
    ],
    "vertical_capabilities": [
         {
            "category": "Vertical",
            "feature_name": "Automated Tax Calculation via HSN Code",
            "functional_specification": "...",
            "non_functional_constraints": "...",
            "source_filename": "example.pdf",
            "section_header": "4.2"
        }
    ],
    "business_logic_and_policies": [
        {
            "rule_title": "Dynamic Price Locking on PR Creation",
            "logic_description": "...",
            "source_filename": "example.pdf",
            "section_header": "4.1"
        }
    ],
    "data_architecture_and_security": [
        "The solution must be deployed on the Adani cloud Platform..."
    ],
    "comprehensive_ecosystem": [
        {
            "system_name": "SAP Ariba Catalogue Buying",
            "integration_purpose": "...",
            "source_filename": "example.pdf",
            "section_header": "1"
        }
    ]
}
"""

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == ".pdf":
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for i in range(min(15, len(reader.pages))): # First 15 pages
                    page_text = reader.pages[i].extract_text()
                    if page_text:
                        text += page_text + "\n"
        elif ext == ".docx":
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs][:200]) # First 200 paragraphs
        elif ext == ".pptx":
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides):
                if i > 15: break
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, read_only=True)
            for sheet in wb.sheetnames[:2]: # First 2 sheets
                ws = wb[sheet]
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 100: break
                    text += " ".join([str(v) for v in row if v is not None]) + "\n"
    except Exception as e:
        print(f"Error reading {file_path}: {e}")
    
    return text[:20000] # Cap text length to avoid token limits

processed_count = 0

for root, _, files in os.walk(rfp_dir):
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        if ext in [".pdf", ".docx", ".pptx", ".xlsx"]:
            if f in processed_files:
                print(f"Skipping {f}, already processed.")
                continue
            file_path = os.path.join(root, f)
            print(f"Processing: {f}")
            text = extract_text(file_path)
            if len(text.strip()) < 50:
                print(f"Skipping {f}, not enough text.")
                continue
                
            prompt = f"""
Analyze the following RFP text and extract the key information into a JSON object matching the schema of the provided example. 
If a specific field is not explicitly mentioned, infer it reasonably or put an empty list/string if inapplicable. Keep it structured.
Return strictly a valid JSON object, no markdown formatting (no ```json ... ```), no other text.

Example JSON Structure:
{example_json}

RFP Text:
{text}
"""
            try:
                response = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {"role": "system", "content": "You are a helpful assistant that extracts structured JSON data from RFPs. You only output valid JSON. Do not include markdown blocks."},
                        {"role": "user", "content": prompt}
                    ],
                    temperature=0.2,
                )
                
                result_text = response.choices[0].message.content.strip()
                if result_text.startswith("```json"):
                    result_text = result_text[7:-3].strip()
                elif result_text.startswith("```"):
                    result_text = result_text[3:-3].strip()
                
                rfp_json = json.loads(result_text)
                rfp_json["_source_file"] = f # Add meta info
                all_rfps.append(rfp_json)
                print(f"Successfully added {f}")
                processed_count += 1
                
                # Save after each successful extraction to not lose progress
                with open(target_json, "w", encoding="utf-8") as out_f:
                    json.dump(all_rfps, out_f, indent=4)
            except Exception as e:
                print(f"Failed to process {f} via API: {e}")

print(f"Finished processing. Added {processed_count} new files.")
